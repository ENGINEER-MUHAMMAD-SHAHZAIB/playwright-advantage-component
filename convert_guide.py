"""
Convert an Advantage Configuration Guide (CSV or Excel) into an automation-ready Excel
for add_components.py.

1. Put the guide file in the  guides  folder.
2. Change the SETTINGS below (file name, program, campus).
3. Run:   python convert_guide.py
4. Open the new file in the data folder, check the "Needs Review" tab.
5. In add_components.py set EXCEL_FILE to the new file.
"""
import csv
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

# ============================ SETTINGS (edit these) ============================
GUIDE_FILE = "guides/Cosmetology - Work Units.csv"   # the guide you received (.csv or .xlsx)
GUIDE_SHEET = None          # for .xlsx guides: sheet name, or None = first sheet
PROGRAM = "Cosmetology"     # program name as it appears after " - " in the course name
CAMPUS = "Royal Treatment Beauty Academy"
STATUS = "Active"
EFFECTIVE_DATE = "01/01/2020"   # grade book "Effective start date" when a new grade book is created
ZERO_WEIGHT_AS = 1              # Workunit Weight 0 or empty in the guide -> use this (None = keep 0)
BLANK_MIN_USE_SUBJECT = True    # Min Workunit Grade empty -> use the subject's Min Grade (e.g. 75)
YN_BLANK_AS = "Yes"             # Required / Must Pass left as "Y/N" in the guide -> "Yes" = switch ON (Advantage default). "No" = OFF, "" = review
OUTPUT_FILE = "data/automation_ready_cosmetology.xlsx"

# Guide "Workunit Type"  ->  Advantage "System component type"
TYPE_MAP = {
    "Written Exam": "Exam",
    "Practical Exam": "Practical Exams",   # Advantage option is "Practical Exams" (from your recording)
    # "Lab Work": "???",
}
PLACEHOLDER_NAMES = {"Exam / Quiz Name"}   # title rows in the guide, not real workunits
# ===============================================================================


def read_guide(path):
    if not os.path.exists(path):
        sys.exit(f"Guide file not found: {path}")
    if path.lower().endswith(".csv"):
        with open(path, encoding="utf-8-sig", newline="") as fh:
            return [list(r) for r in csv.reader(fh)]
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[GUIDE_SHEET] if GUIDE_SHEET else wb.worksheets[0]
    return [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]


def find_columns(rows):
    """Find the header row and the column of each field (left-hand table only)."""
    wanted = {
        "subject": "subject name", "subject_weight": "subject weight",
        "subject_min": "min subject grade", "seq": "seq", "name": "workunit name",
        "type": "workunit type", "weight": "workunit weight", "min": "min workunit grade",
        "lab": "lab count", "required": "required", "must_pass": "must pass",
    }
    for i, row in enumerate(rows):
        low = [c.strip().lower() for c in row]
        if any(c.startswith("subject name") for c in low) and any(c.startswith("workunit name") for c in low):
            cols = {}
            for key, text in wanted.items():
                for j, c in enumerate(low):
                    if c.startswith(text) and key not in cols:
                        cols[key] = j
                        break
            missing = [k for k in wanted if k not in cols]
            if missing:
                sys.exit(f"Header row found but these columns are missing: {missing}")
            subj_label = cols["subject"] - 1   # "Subject 01" label sits left of Subject Name
            return i, cols, subj_label
    sys.exit("Could not find the header row (needs 'Subject Name' and 'Workunit Name').")


def to_num(v):
    v = str(v).strip()
    try:
        f = float(v)
        return int(f) if f.is_integer() else f
    except ValueError:
        return v


def parse(rows):
    hdr, c, subj_label = find_columns(rows)
    comps, subjects, review, shifted = [], [], [], set()
    subj, subj_min = None, ""
    for r in rows[hdr + 1:]:
        r = r + [""] * (max(c.values()) + 2 - len(r))
        get = lambda k: str(r[c[k]]).strip()
        label = str(r[subj_label]).strip() if subj_label >= 0 else ""
        if get("subject") and (label.lower().startswith("subject") or not label):
            subj = get("subject")
            subj_min = to_num(get("subject_min"))
            subjects.append([label or f"Subject {len(subjects) + 1:02d}", subj,
                             to_num(get("subject_weight")), to_num(get("subject_min"))])
        if not subj:
            continue
        name, seq = get("name"), get("seq")
        if not name:
            continue
        if name in PLACEHOLDER_NAMES:
            shifted.add(subj)
            review.append([subj, seq, name, "Skipped: title row. Seq for this subject renumbered from 1"])
            continue
        if subj in shifted and seq.isdigit():
            seq = str(int(seq) - 1)
        typ = get("type")
        adv_type = TYPE_MAP.get(typ, "")
        req, mp = get("required"), get("must_pass")
        if YN_BLANK_AS and req not in ("Yes", "No"):
            req = YN_BLANK_AS
        if YN_BLANK_AS and mp not in ("Yes", "No"):
            mp = YN_BLANK_AS
        notes, run = [], "Yes"
        if not adv_type:
            notes.append(f"Advantage type for '{typ}' not known - add it to TYPE_MAP"); run = "No"
        if req not in ("Yes", "No") or mp not in ("Yes", "No"):
            notes.append("Required / Must Pass not filled (Y/N)"); run = "No"
        lab = get("lab")
        weight, minv = to_num(get("weight")), to_num(get("min"))
        if ZERO_WEIGHT_AS is not None and weight in ("", 0):
            weight = ZERO_WEIGHT_AS
        if BLANK_MIN_USE_SUBJECT and minv == "":
            minv = subj_min
        comps.append([subj, f"{subj} - {PROGRAM}", to_num(seq), name, typ, adv_type,
                      to_num(lab) if lab else 0, weight, minv,
                      req, mp, run, "; ".join(notes), f"{subj} [{PROGRAM}]", subj_min])
        if notes:
            review.append([subj, seq, name, "; ".join(notes)])
    for s in subjects:
        if not any(x[0] == s[1] for x in comps):
            review.append([s[1], "", "", "Skipped: subject has no workunits in the guide"])
    return comps, subjects, review


def write(comps, subjects, review):
    F = "Arial"
    hdr_fill = PatternFill("solid", fgColor="1F4E79")
    hdr_font = Font(name=F, bold=True, color="FFFFFF")
    body = Font(name=F, size=10)
    edit = PatternFill("solid", fgColor="FFF9E6")

    def header(ws, n):
        for i in range(1, n + 1):
            cell = ws.cell(1, i)
            cell.fill, cell.font = hdr_fill, hdr_font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 32

    wb = Workbook()
    ws = wb.active
    ws.title = "Components"
    H = ["Row ID", "Run", "Subject", "Course (open in Advantage)", "Seq", "Component Name (Description)",
         "Name Length", "Over 25 chars?", "Type in Guide", "System Component Type", "Status", "Campus",
         "Number", "Weight", "Min Score", "Required", "Must Pass", "Result", "Notes",
         "Course Name (list)", "Course Min Score", "Grade Book Effective Date"]
    ws.append(H)
    for i, (subj, course, seq, name, typ, atype, lab, w, mn, req, mp, run, notes, clist, cmin) in enumerate(comps, 1):
        r = i + 1
        ws.append([i, run, subj, course, seq, name, f"=LEN(F{r})", f'=IF(G{r}>25,"Yes","No")',
                   typ, atype, STATUS, CAMPUS, lab, w, mn, req, mp, "", notes, clist, cmin, EFFECTIVE_DATE])
    header(ws, len(H))
    for i, w in enumerate([7, 6, 24, 40, 6, 52, 8, 9, 14, 14, 9, 30, 8, 8, 9, 9, 9, 12, 44, 40, 10, 14], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    last = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=last):
        for cell in row:
            cell.font = body
            cell.border = Border(bottom=Side(style="thin", color="D0D0D0"))
    for col in ["B", "J", "K", "L", "M", "N", "O", "P", "Q", "T", "U", "V"]:
        for r in range(2, last + 1):
            ws[f"{col}{r}"].fill = edit
    dv = DataValidation(type="list", formula1='"Yes,No"')
    ws.add_data_validation(dv)
    dv.add(f"B2:B{last}")
    dv.add(f"P2:Q{last}")
    ws.conditional_formatting.add(f"H2:H{last}", FormulaRule(formula=['$H2="Yes"'], fill=PatternFill("solid", fgColor="FDE2E1")))
    ws.conditional_formatting.add(f"B2:B{last}", FormulaRule(formula=['$B2="No"'], fill=PatternFill("solid", fgColor="E0E0E0")))
    ws.auto_filter.ref = f"A1:V{last}"
    ws["D1"].comment = Comment("Grade book name in Advantage ('Subject - Program'). Created if missing.", "tool")
    ws["T1"].comment = Comment("Course name in the ADD/VIEW COURSE list ('Subject [Program]'). Created if missing.", "tool")
    ws["J1"].comment = Comment("Value picked in Advantage's 'System component type'. Comes from TYPE_MAP in convert_guide.py.", "tool")

    s2 = wb.create_sheet("Subjects")
    s2.append(["Subject #", "Subject Name", "Subject Weight", "Min Subject Grade %", "Course name in Advantage", "Components"])
    for s in subjects:
        r = s2.max_row + 1
        s2.append(s + [f"{s[1]} - {PROGRAM}", f"=COUNTIF(Components!C:C,B{r})"])
    header(s2, 6)

    s3 = wb.create_sheet("Needs Review")
    s3.append(["Subject", "Seq", "Name", "Why"])
    for r in review:
        s3.append(r)
    header(s3, 4)
    for sh, widths in ((s2, [11, 32, 14, 18, 44, 12]), (s3, [28, 6, 44, 70])):
        for i, w in enumerate(widths, 1):
            sh.column_dimensions[sh.cell(1, i).column_letter].width = w
        for row in sh.iter_rows(min_row=2):
            for cell in row:
                cell.font = body

    os.makedirs(os.path.dirname(OUTPUT_FILE) or ".", exist_ok=True)
    if os.path.exists(OUTPUT_FILE):
        sys.exit(f"{OUTPUT_FILE} already exists - choose a new OUTPUT_FILE name (or delete the old one).")
    wb.save(OUTPUT_FILE)


if __name__ == "__main__":
    comps, subjects, review = parse(read_guide(GUIDE_FILE))
    write(comps, subjects, review)
    ready = sum(1 for x in comps if x[11] == "Yes")
    print(f"Saved {OUTPUT_FILE}")
    print(f"  Subjects: {len(subjects)}   Components: {len(comps)}   Run = Yes: {ready}   Run = No: {len(comps) - ready}")
    print(f"  Names over 25 characters: {sum(1 for x in comps if len(x[3]) > 25)}")
    print(f"  Rows to review: {len(review)}  (see the 'Needs Review' tab)")
    print(f"Next: in add_components.py set  EXCEL_FILE = \"{OUTPUT_FILE}\"")
